from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from app.ingestion.loader import (
    DocumentLoadError,
    load_document,
)


def test_load_txt_document(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "sample.txt"

    file_path.write_text(
        "Clinical guideline sample text.",
        encoding="utf-8",
    )

    document = load_document(file_path)

    assert document.filename == "sample.txt"
    assert document.file_type == "txt"
    assert document.full_text == (
        "Clinical guideline sample text."
    )
    assert document.character_count > 0
    assert len(document.pages) == 1
    assert document.pages[0].page_number is None


def test_load_docx_document(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "sample.docx"

    source_document = DocxDocument()
    source_document.add_heading(
        "Clinical Guidance",
        level=1,
    )
    source_document.add_paragraph(
        "This is a test paragraph."
    )
    source_document.save(file_path)

    document = load_document(file_path)

    assert document.file_type == "docx"
    assert "Clinical Guidance" in document.full_text
    assert "This is a test paragraph." in document.full_text
    assert len(document.pages) == 1
    assert document.pages[0].page_number is None


def test_empty_txt_document_is_rejected(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "empty.txt"

    file_path.write_text(
        "",
        encoding="utf-8",
    )

    with pytest.raises(
        DocumentLoadError,
        match="Document is empty",
    ):
        load_document(file_path)


def test_missing_document_is_rejected(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "missing.pdf"

    with pytest.raises(
        DocumentLoadError,
        match="does not exist",
    ):
        load_document(file_path)


def test_unsupported_file_type_is_rejected(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "sample.csv"

    file_path.write_text(
        "a,b,c",
        encoding="utf-8",
    )

    with pytest.raises(
        DocumentLoadError,
        match="Unsupported file type",
    ):
        load_document(file_path)


def test_pdf_with_no_extractable_text_is_rejected(
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "blank.pdf"

    writer = PdfWriter()
    writer.add_blank_page(
        width=612,
        height=792,
    )

    with file_path.open("wb") as output_file:
        writer.write(output_file)

    with pytest.raises(
        DocumentLoadError,
        match="No readable text was extracted",
    ):
        load_document(file_path)