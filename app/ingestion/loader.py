from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.models.documents import (
    DocumentPage,
    LoadedDocument,
    SupportedFileType,
)


class DocumentLoadError(RuntimeError):
    """
    Raised when a supported document cannot be safely loaded.
    """


SUPPORTED_EXTENSIONS: dict[str, SupportedFileType] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "txt",
}


def load_document(file_path: str | Path) -> LoadedDocument:
    """
    Load a supported document and convert it into the common
    LoadedDocument representation.
    """
    path = Path(file_path)

    _validate_source_file(path)

    suffix = path.suffix.lower()
    file_type = SUPPORTED_EXTENSIONS[suffix]

    try:
        if file_type == "pdf":
            pages = _load_pdf(path)

        elif file_type == "docx":
            pages = _load_docx(path)

        elif file_type == "txt":
            pages = _load_txt(path)

        else:
            raise DocumentLoadError(
                f"Unsupported document type: {suffix}"
            )

    except DocumentLoadError:
        raise

    except Exception as exc:
        raise DocumentLoadError(
            f"Failed to load document: {path.name}"
        ) from exc

    usable_pages = [
        page
        for page in pages
        if page.text.strip()
    ]

    if not usable_pages:
        raise DocumentLoadError(
            f"No readable text was extracted from: {path.name}"
        )

    full_text = "\n\n".join(
        page.text.strip()
        for page in usable_pages
    ).strip()

    if not full_text:
        raise DocumentLoadError(
            f"Extracted document text is empty: {path.name}"
        )

    return LoadedDocument(
        filename=path.name,
        file_path=path.resolve(),
        file_type=file_type,
        pages=usable_pages,
        full_text=full_text,
        character_count=len(full_text),
    )


def _validate_source_file(path: Path) -> None:
    """
    Validate a source file before attempting extraction.
    """
    if not path.exists():
        raise DocumentLoadError(
            f"Document does not exist: {path}"
        )

    if not path.is_file():
        raise DocumentLoadError(
            f"Path is not a file: {path}"
        )

    suffix = path.suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(
            sorted(SUPPORTED_EXTENSIONS)
        )

        raise DocumentLoadError(
            f"Unsupported file type '{suffix}'. "
            f"Supported types are: {supported}"
        )

    if path.stat().st_size == 0:
        raise DocumentLoadError(
            f"Document is empty: {path.name}"
        )


def _load_pdf(path: Path) -> list[DocumentPage]:
    """
    Extract text page-by-page from a PDF.
    """
    reader = PdfReader(str(path))

    if not reader.pages:
        raise DocumentLoadError(
            f"PDF contains no pages: {path.name}"
        )

    pages: list[DocumentPage] = []

    for index, pdf_page in enumerate(
        reader.pages,
        start=1,
    ):
        text = pdf_page.extract_text() or ""

        if text.strip():
            pages.append(
                DocumentPage(
                    page_number=index,
                    text=text.strip(),
                )
            )

    return pages


def _load_docx(path: Path) -> list[DocumentPage]:
    """
    Extract paragraphs and useful table text from a DOCX file.

    DOCX files do not provide reliable physical page boundaries,
    so the result is represented as one logical text unit.
    """
    document = DocxDocument(str(path))

    blocks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            style_name = (
                paragraph.style.name
                if paragraph.style is not None
                else ""
            )

            if style_name.startswith("Heading"):
                blocks.append(f"\n{text}\n")
            else:
                blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if cells:
                blocks.append(" | ".join(cells))

    text = "\n".join(blocks).strip()

    if not text:
        return []

    return [
        DocumentPage(
            page_number=None,
            text=text,
        )
    ]


def _load_txt(path: Path) -> list[DocumentPage]:
    """
    Load a UTF-8 plain-text document.
    """
    try:
        text = path.read_text(
            encoding="utf-8",
        )

    except UnicodeDecodeError as exc:
        raise DocumentLoadError(
            f"TXT document is not valid UTF-8: {path.name}"
        ) from exc

    text = text.strip()

    if not text:
        return []

    return [
        DocumentPage(
            page_number=None,
            text=text,
        )
    ]